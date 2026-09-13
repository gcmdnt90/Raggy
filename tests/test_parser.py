"""Tests for the Raggy document parser."""


import pytest

from app.parsers.documents import extract_text, list_supported_files, load_document


class TestExtractText:
    """Test document text extraction."""

    def test_extract_md(self, tmp_path):
        md_file = tmp_path / "test.md"
        md_file.write_text("# Title\n\nSome content here.", encoding="utf-8")
        text = extract_text(md_file)
        assert "# Title" in text
        assert "Some content here." in text

    def test_extract_txt(self, tmp_path):
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("Plain text content.", encoding="utf-8")
        text = extract_text(txt_file)
        assert text == "Plain text content."

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            extract_text("/nonexistent/path.md")

    def test_unsupported_format(self, tmp_path):
        docx = tmp_path / "test.rtf"
        docx.write_text("fake", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            extract_text(docx)

    def test_rejects_invalid_docx_magic(self, tmp_path):
        docx = tmp_path / "test.docx"
        docx.write_bytes(b"fake")
        with pytest.raises(ValueError, match="magic"):
            load_document(docx)

    def test_extract_docx(self, tmp_path):
        pytest.importorskip("docx")
        from docx import Document

        docx = tmp_path / "test.docx"
        document = Document()
        document.add_paragraph("Docx content")
        document.save(docx)
        assert "Docx content" in extract_text(docx)

    def test_extract_xlsx(self, tmp_path):
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx = tmp_path / "test.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "Xlsx content"
        workbook.save(xlsx)
        workbook.close()
        assert "Xlsx content" in extract_text(xlsx)


class TestListSupportedFiles:
    """Test file discovery."""

    def test_finds_supported_files(self, tmp_path):
        (tmp_path / "doc.md").write_text("md", encoding="utf-8")
        (tmp_path / "note.txt").write_text("txt", encoding="utf-8")
        (tmp_path / "doc.docx").write_bytes(b"PK\x03\x04fake")
        (tmp_path / "sheet.xlsx").write_bytes(b"PK\x03\x04fake")
        (tmp_path / "skip.rtf").write_text("no", encoding="utf-8")

        files = list_supported_files(tmp_path)
        names = [f.name for f in files]
        assert "doc.md" in names
        assert "note.txt" in names
        assert "doc.docx" in names
        assert "sheet.xlsx" in names
        assert "skip.rtf" not in names

    def test_recursive(self, tmp_path):
        sub = tmp_path / "subdir"
        sub.mkdir()
        (sub / "nested.md").write_text("nested", encoding="utf-8")

        files = list_supported_files(tmp_path)
        assert any(f.name == "nested.md" for f in files)

    def test_empty_dir(self, tmp_path):
        assert list_supported_files(tmp_path) == []

    def test_nonexistent_dir(self):
        assert list_supported_files("/nonexistent") == []
