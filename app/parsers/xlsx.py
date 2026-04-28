"""XLSX document loader."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from app.parsers.documents import ParsedDocument
from app.utils.sanitize import XLSX_MAX_CELLS, sanitize_text, validate_xlsx


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return sanitize_text(str(value))


def load_xlsx(path: str | Path) -> ParsedDocument:
    """Load and validate an XLSX workbook in read-only/data-only mode."""
    p = Path(path)
    if p.suffix.lower() in {".xlsm", ".xlsb"}:
        raise ValueError("Macro or binary Excel workbooks are not supported.")

    validated = validate_xlsx(p)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required for XLSX parsing.") from exc

    workbook = load_workbook(
        io.BytesIO(validated.data),
        read_only=True,
        data_only=True,
    )
    cell_count = 0
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"## Sheet: {sanitize_text(sheet.title)}")
            for row in sheet.iter_rows(values_only=True):
                cell_count += len(row)
                if cell_count > XLSX_MAX_CELLS:
                    raise ValueError(
                        f"XLSX has too many cells: {cell_count} exceeds {XLSX_MAX_CELLS}."
                    )
                values = [_cell_to_text(value).strip() for value in row]
                values = [value for value in values if value]
                if values:
                    parts.append("\t".join(values))
    finally:
        workbook.close()

    return ParsedDocument(
        text="\n".join(parts),
        metadata={
            "source_file": str(p),
            "format": "xlsx",
            "size_bytes": validated.size_bytes,
            "cell_count": cell_count,
        },
    )
