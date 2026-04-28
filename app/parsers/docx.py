"""DOCX document loader."""

from __future__ import annotations

import io
from pathlib import Path

from app.parsers.documents import ParsedDocument
from app.utils.sanitize import sanitize_text, validate_docx


def load_docx(path: str | Path) -> ParsedDocument:
    """Load and validate a DOCX document."""
    p = Path(path)
    if p.suffix.lower() == ".docm":
        raise ValueError("Macro-enabled Word documents (.docm) are not supported.")

    validated = validate_docx(p)
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required for DOCX parsing.") from exc

    document = Document(io.BytesIO(validated.data))
    parts = [sanitize_text(paragraph.text) for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            cells = [sanitize_text(cell.text).strip() for cell in row.cells if cell.text]
            if cells:
                parts.append("\t".join(cells))

    return ParsedDocument(
        text="\n".join(parts),
        metadata={
            "source_file": str(p),
            "format": "docx",
            "size_bytes": validated.size_bytes,
        },
    )
